import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, name='Times New Roman', size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic

def add_heading(doc, text, level=1, centered=False):
    p = doc.add_paragraph()
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    if level == 0: # Chapter
        set_font(run, size=16, bold=True)
    elif level == 1: # Section
        set_font(run, size=14, bold=True)
    else: # Subsection
        set_font(run, size=12, bold=True)
    return p

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        set_font(run, size=12)
    return p

def create_report():
    doc = Document()
    
    # Title Page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n\n\n\nECOMMERCE STORE FOR RETAIL BUSINESS OWNERS\n\n")
    set_font(run, size=22, bold=True)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("A Project Report Submitted in partial fulfillment of the requirements for the degree of\nBACHELOR OF ENGINEERING\nin\nCOMPUTER SCIENCE AND ENGINEERING")
    set_font(run, size=12)

    doc.add_page_break()

    # Bonafide Certificate
    add_heading(doc, "BONAFIDE CERTIFICATE", level=0, centered=True)
    add_body(doc, "\nCertified that this project report titled \"ECOMMERCE STORE FOR RETAIL BUSINESS OWNERS\" is the bonafide work of HARI HARA SUDHAN (Admin) who carried out the project work under my supervision.\n\n\n\n")
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run1 = p.add_run("SIGNATURE\t\t\t\t\tSIGNATURE\n")
    run2 = p.add_run("HEAD OF THE DEPARTMENT\t\t\tPROJECT GUIDE")
    set_font(run1, bold=True)
    set_font(run2, bold=True)

    doc.add_page_break()

    # Acknowledgement
    add_heading(doc, "ACKNOWLEDGEMENT", level=0, centered=True)
    add_body(doc, "I wish to express my sincere gratitude to the management and administration for providing the necessary infrastructure and environment for this project. I am deeply thankful to the Head of the Department and my Project Guide for their constant encouragement and valuable guidance throughout the course of this project. Finally, I would like to thank my peers and family for their unwavering support.")

    doc.add_page_break()

    # Abstract
    add_heading(doc, "ABSTRACT", level=0, centered=True)
    add_body(doc, "This project involves the design and implementation of a premium e-commerce platform specifically tailored for retail business owners. The objective was to transform a generic online marketplace into a high-end, brand-focused shopping experience inspired by industry leaders like Gymshark and Shopify. The application utilizes a robust stack comprising Node.js for the backend, MySQL for data persistence, and native CSS/JS for a high-performance, minimalist frontend. Key features include a Gymshark-inspired product detail system with interactive accordions, a Shopify-style analytics dashboard for store owners, and a premium design system featuring a pearl-grey and slate color palette. The resulting system provides an elite-tier boutique experience that maximizes user engagement and conversion rates.")

    doc.add_page_break()

    # Table of Contents (Placeholder)
    add_heading(doc, "TABLE OF CONTENTS", level=0, centered=True)
    doc.add_paragraph("Chapter 1 Introduction\nChapter 2 Sustainability Goals\nChapter 3 System Requirements Specification\nChapter 4 System Design\nChapter 5 Implementation\nChapter 6 Results and Discussion\nChapter 7 Conclusion")

    doc.add_page_break()

    # Chapter 1
    add_heading(doc, "CHAPTER 1\nINTRODUCTION", level=0, centered=True)
    add_heading(doc, "1.1 Background", level=1)
    add_body(doc, "In the modern retail landscape, an online presence is no longer optional. However, most small to medium retail business owners are stuck with generic, uninspiring templates that fail to convey the quality of their brand. This project, RetailMart, was conceived to bridge the gap between expensive custom-built agencies and low-end generic templates.")
    
    add_heading(doc, "1.2 Objectives", level=1)
    add_body(doc, "The primary objective is to build a high-performance, visually stunning e-commerce store that feels like a premium brand. Specific goals include implementing a sophisticated design system, interactive product features, and a comprehensive data-driven dashboard for owners.")

    # Chapter 2
    doc.add_page_break()
    add_heading(doc, "CHAPTER 2\nSUSTAINABILITY GOALS", level=0, centered=True)
    add_body(doc, "This project aligns with UN Sustainable Development Goal 8: Decent Work and Economic Growth by empowering local retail owners with professional digital tools to grow their businesses. It also supports SDG 9: Industry, Innovation and Infrastructure by utilizing modern web technologies to create resilient and accessible digital infrastructure.")

    # Chapter 3
    doc.add_page_break()
    add_heading(doc, "CHAPTER 3\nSYSTEM REQUIREMENTS SPECIFICATION", level=0, centered=True)
    add_heading(doc, "3.1 Hardware Requirements", level=1)
    add_body(doc, "Processor: Intel Core i5 or equivalent\nRAM: 8GB minimum\nStorage: 256GB SSD")
    add_heading(doc, "3.2 Software Requirements", level=1)
    add_body(doc, "Operating System: Windows 10/11\nEnvironment: Node.js v20+\nDatabase: MySQL 8.0\nEditor: VS Code")

    # Chapter 4
    doc.add_page_break()
    add_heading(doc, "CHAPTER 4\nSYSTEM DESIGN", level=0, centered=True)
    add_body(doc, "The system follows the Model-View-Controller (MVC) architectural pattern. The backend is built using Express.js, handling routing and database interactions. The frontend is a Single Page Application (SPA) style implementation using native JS and CSS for maximum speed and control.")

    # Chapter 5
    doc.add_page_break()
    add_heading(doc, "CHAPTER 5\nIMPLEMENTATION", level=0, centered=True)
    add_body(doc, "The implementation phase involved creating a custom MySQL schema to support product variations (colors, sizes) and analytics tracking. The UI was built using a custom design system with CSS variables for consistency. The product detail page features an interactive accordion system for materials, size fit, and delivery information.")

    # Chapter 6
    doc.add_page_break()
    add_heading(doc, "CHAPTER 6\nRESULTS AND DISCUSSION", level=0, centered=True)
    add_body(doc, "The final application demonstrates a significant improvement in UI/UX over traditional templates. The Shopify-style admin dashboard provides real-time insights into sales trends and customer behavior, while the Gymshark-inspired product pages provide a highly engaging shopping experience.")

    # Chapter 7
    doc.add_page_break()
    add_heading(doc, "CHAPTER 7\nCONCLUSION", level=0, centered=True)
    add_body(doc, "The Ecommerce Store for Retail Business Owners project successfully met all its design and functional goals. It provides a blueprint for how retail owners can leverage modern web design and data analytics to compete in a crowded digital marketplace.")

    output_path = os.path.join("c:\\Users\\Admin\\Videos\\FS_PROJECT", "Project_Report_RetailMart.docx")
    doc.save(output_path)
    print(f"Report saved to {output_path}")

if __name__ == "__main__":
    create_report()
